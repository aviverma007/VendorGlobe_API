"""Collection dashboard data — read the CRM team's three Excel files
straight from the shared folder, no database needed.

Folder (override with VG_COLLECTION_DIR):
    \\\\WIN-PJQA0USC6HT\\Users\\anirudh.verma\\Downloads\\CRM\\CRM DATA REPORTS\\CRM\\COLLECTION
Files (fixed names; falls back to the newest file matching the prefix,
so dated names like "PTP month of Sep-26.xlsx" also work):
    Collection Master.xlsx   -> PDC sheet (post-dated cheques in hand)
    Daily Collection Report.xlsx -> one receipts sheet per project
    PTP.xlsx                 -> Sheet1, the full per-unit customer ledger

The endpoint caches each parsed file keyed by its modified-time: the
team saves the Excel, the next dashboard load re-reads only what
changed. Files are copied to a temp path before opening so a workbook
someone has open in Excel never breaks the API. Excel error strings
(#N/A, #REF!, ...) and the pre-filled empty rows are dropped.
"""
import os
import glob
import shutil
import tempfile
import datetime
import threading

import openpyxl
from flask import jsonify

COLLECTION_DIR = os.environ.get(
    "VG_COLLECTION_DIR",
    r"\\WIN-PJQA0USC6HT\Users\anirudh.verma\Downloads\CRM\CRM DATA REPORTS\CRM\COLLECTION",
)
FILES = {
    "master": ("Collection Master.xlsx", "Collection Master*"),
    "daily": ("Daily Collection Report.xlsx", "Daily Collection*"),
    "ptp": ("PTP.xlsx", "PTP*"),
}
# Daily workbook: sheets that are pivots/rosters, not receipt logs
DAILY_SKIP = {"summary", "rm sheet", "final", "fina l-2", "aug - aop sheet", "m3m",
              "september roaster"}

_lock = threading.Lock()
_cache = {}          # key -> (mtime, parsed)

ERRS = {"#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "0", "", "None"}


def _clean(v):
    if v is None:
        return None
    s = str(v).strip()
    return None if s in ("", "#N/A", "#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "None") else s


def _num(v):
    try:
        f = float(v)
        return round(f, 2)
    except (TypeError, ValueError):
        return 0.0


def _iso(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    s = _clean(v)
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d-%b-%Y", "%d-%b-%y"):
        try:
            return datetime.datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _norm_proj(p):
    s = (_clean(p) or "").upper()
    s = " ".join(s.split())
    s = s.replace(" - ", "-").replace(" -", "-").replace("- ", "-")
    return s


def _find(kind):
    fixed, pattern = FILES[kind]
    path = os.path.join(COLLECTION_DIR, fixed)
    if os.path.exists(path):
        return path
    hits = sorted(glob.glob(os.path.join(COLLECTION_DIR, pattern + ".xlsx")),
                  key=os.path.getmtime, reverse=True)
    return hits[0] if hits else None


def _open_copy(path):
    """Copy to temp first so an Excel lock never breaks the read."""
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    shutil.copyfile(path, tmp)
    try:
        return openpyxl.load_workbook(tmp, read_only=True, data_only=True), tmp
    except Exception:
        os.unlink(tmp)
        raise


def _parse_ptp(path):
    wb, tmp = _open_copy(path)
    try:
        ws = wb["Sheet1"]
        rows = ws.iter_rows(values_only=True)
        next(rows)  # header
        out = []
        for r in rows:
            if len(r) < 46 or not _clean(r[3]):     # Reg.No required
                continue
            out.append({
                "proj": _norm_proj(r[1]), "phase": _clean(r[2]), "reg": _clean(r[3]),
                "unit": _clean(r[4]), "allot": _clean(r[5]), "name": _clean(r[6]),
                "profile": _clean(r[7]), "allotDate": _iso(r[8]), "unitType": _clean(r[9]),
                "plan": _clean(r[10]), "planType": _clean(r[11]), "broker": _clean(r[12]),
                "area": _num(r[15]), "rate": _num(r[17]),
                "tcv": _num(r[18]), "dem": _num(r[19]), "rec": _num(r[20]), "due": _num(r[21]),
                "recPct": _num(r[24]),
                "letter": _clean(r[25]), "letterDate": _iso(r[26]), "letterDue": _iso(r[27]),
                "rm": _clean(r[28]) or _clean(r[29]),
                "remarks": (_clean(r[30]) or "")[:220] or None,
                "rmStatus": _clean(r[31]),
                "funding": _clean(r[34]), "bank": _clean(r[35]),
                "sanctDate": _iso(r[36]), "sanctAmt": _num(r[38]),
                "bba": _clean(r[39]), "bbaDate": _iso(r[40]),
                "ptpDate": _iso(r[45]),
                "possession": _clean(r[47]) if len(r) > 47 else None,
            })
        return out
    finally:
        wb.close()
        os.unlink(tmp)


def _parse_daily(path):
    wb, tmp = _open_copy(path)
    try:
        out = []
        for ws in wb.worksheets:
            if ws.title.strip().lower() in DAILY_SKIP:
                continue
            rows = ws.iter_rows(values_only=True)
            try:
                next(rows)  # header
            except StopIteration:
                continue
            proj = _norm_proj(ws.title)
            for r in rows:
                if len(r) < 13:
                    continue
                amt = _num(r[4])
                if amt == 0 and not _iso(r[8]):
                    continue                      # pre-filled empty row
                out.append({
                    "proj": proj, "reg": _clean(r[1]), "name": _clean(r[2]),
                    "unit": _clean(r[3]), "amt": amt, "mode": _clean(r[5]),
                    "chq": _clean(r[6]), "bank": _clean(r[7]),
                    "rcptDate": _iso(r[8]) or _iso(r[9]) or _iso(r[11]), "chqDate": _iso(r[9]),
                    "clearDate": _iso(r[10]), "created": _clean(r[11]),
                    "rm": _clean(r[12]),
                    "milestone": _clean(r[15]) if len(r) > 15 else None,
                    "dueDate": _iso(r[16]) if len(r) > 16 else None,
                })
        return out
    finally:
        wb.close()
        os.unlink(tmp)


def _parse_master(path):
    wb, tmp = _open_copy(path)
    try:
        out = []
        if "PDC" in wb.sheetnames:
            ws = wb["PDC"]
            rows = ws.iter_rows(values_only=True)
            next(rows, None)          # blank/ratio row
            next(rows, None)          # header row
            for r in rows:
                if len(r) < 10:
                    continue
                amt = _num(r[9])
                if amt == 0 or not _clean(r[2]):
                    continue
                out.append({
                    "proj": _norm_proj(r[1]), "reg": _clean(r[2]), "given": _iso(r[3]),
                    "unit": _clean(r[4]), "mode": _clean(r[5]), "chq": _clean(r[6]),
                    "chqDate": _iso(r[7]), "bank": _clean(r[8]), "amt": amt,
                    "allotDate": _iso(r[10]), "phase": _clean(r[11]),
                    "tower": _clean(r[12]), "received": _iso(r[13]),
                })
        return out
    finally:
        wb.close()
        os.unlink(tmp)


PARSERS = {"ptp": _parse_ptp, "daily": _parse_daily, "master": _parse_master}


def _load(kind):
    path = _find(kind)
    if not path:
        return None, None, f"file not found for '{kind}' in {COLLECTION_DIR}"
    mtime = os.path.getmtime(path)
    with _lock:
        hit = _cache.get(kind)
        if hit and hit[0] == mtime:
            return hit[1], mtime, None
    data = PARSERS[kind](path)
    with _lock:
        _cache[kind] = (mtime, data)
    return data, mtime, None


def register(app):
    @app.route("/collections/data")
    def collections_data():
        try:
            ledger, m1, e1 = _load("ptp")
            receipts, m2, e2 = _load("daily")
            pdc, m3, e3 = _load("master")
            errors = [e for e in (e1, e2, e3) if e]

            def pack(rows):
                if not rows:
                    return {"cols": [], "rows": []}
                cols = list(rows[0].keys())
                return {"cols": cols, "rows": [[r.get(c) for c in cols] for r in rows]}

            return jsonify({
                "ok": True,
                "asOf": {
                    "ptp": datetime.datetime.fromtimestamp(m1).isoformat() if m1 else None,
                    "daily": datetime.datetime.fromtimestamp(m2).isoformat() if m2 else None,
                    "master": datetime.datetime.fromtimestamp(m3).isoformat() if m3 else None,
                },
                "errors": errors,
                "ledger": pack(ledger),
                "receipts": pack(receipts),
                "pdc": pack(pdc),
            })
        except Exception as exc:            # noqa: BLE001 — surface to the dashboard
            return jsonify({"ok": False, "error": str(exc)}), 500

    @app.route("/collections/health")
    def collections_health():
        out = {}
        for kind in FILES:
            path = _find(kind)
            out[kind] = {
                "path": path,
                "exists": bool(path and os.path.exists(path)),
                "modified": datetime.datetime.fromtimestamp(os.path.getmtime(path)).isoformat()
                if path and os.path.exists(path) else None,
            }
        return jsonify({"ok": True, "dir": COLLECTION_DIR, "files": out})
